"""Native Requirements Document Reader (DocReader).

Zero-friction, zero-external-dependency parser for technical requirement documents:
- .docx (via pure zipfile + ElementTree XML parser, with python-docx if installed)
- .txt, .md, .markdown, .rst
- .json

Enables AI agents to immediately extract text, structure, and tables from user-provided
documents without running terminal commands, installing packages, or requesting approvals.
"""

import json
import os
import xml.etree.ElementTree as ET
import zipfile
from typing import Optional


class RequirementsDocReader:
    """Universal document reader for technical specifications (ТЗ)."""

    @staticmethod
    def read_document(filepath: str, max_chars: int = 100_000) -> str:
        """Read and format requirements document content.

        Args:
            filepath: Path to the target document (.docx, .txt, .md, .json).
            max_chars: Maximum characters to return to prevent overflowing context.

        Returns:
            Extracted text formatted cleanly in Markdown.
        """
        clean_path = os.path.abspath(filepath.strip())
        if not os.path.exists(clean_path):
            raise FileNotFoundError(f"Requirements file not found at: {clean_path}")

        ext = os.path.splitext(clean_path)[1].lower()

        if ext == ".docx":
            content = RequirementsDocReader._read_docx(clean_path)
        elif ext in (".txt", ".md", ".markdown", ".rst"):
            content = RequirementsDocReader._read_text(clean_path)
        elif ext == ".json":
            content = RequirementsDocReader._read_json(clean_path)
        else:
            # Fallback: attempt binary zip check or text read
            try:
                content = RequirementsDocReader._read_docx(clean_path)
            except Exception:
                content = RequirementsDocReader._read_text(clean_path)

        if len(content) > max_chars:
            truncated = content[:max_chars]
            return f"{truncated}\n\n... [TRUNCATED: document exceeded {max_chars} characters] ..."
        return content

    @staticmethod
    def _read_docx(filepath: str) -> str:
        """Parse .docx file. Tries python-docx first; falls back to pure standard library."""
        # Method 1: python-docx if available
        try:
            import docx

            doc = docx.Document(filepath)
            lines = []
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    style_name = (p.style.name if p.style else "").lower()
                    if "heading 1" in style_name:
                        lines.append(f"\n# {text}\n")
                    elif "heading 2" in style_name:
                        lines.append(f"\n## {text}\n")
                    elif "heading 3" in style_name:
                        lines.append(f"\n### {text}\n")
                    else:
                        lines.append(text)

            for table in doc.tables:
                lines.append("\n")
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                lines.append("| " + " | ".join(headers) + " |")
                lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for row in table.rows[1:]:
                    vals = [cell.text.strip() for cell in row.cells]
                    lines.append("| " + " | ".join(vals) + " |")
                lines.append("\n")

            result = "\n".join(lines).strip()
            if result:
                return result
        except ImportError:
            pass
        except Exception:
            pass

        # Method 2: Pure Python standard library (zipfile + XML)
        # .docx is a zip archive containing word/document.xml
        try:
            with zipfile.ZipFile(filepath, "r") as z:
                xml_content = z.read("word/document.xml")
            
            tree = ET.fromstring(xml_content)
            # Word namespaces
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

            lines = []
            # Extract paragraphs and tables in order
            body = tree.find("w:body", ns)
            if body is None:
                body = tree

            for child in body:
                tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

                if tag == "p":
                    # Paragraph
                    texts = [t.text for t in child.findall(".//w:t", ns) if t.text]
                    p_text = "".join(texts).strip()
                    if p_text:
                        # Check heading style
                        p_style = child.find(".//w:pStyle", ns)
                        val = p_style.attrib.get(f"{{{ns['w']}}}val", "") if p_style is not None else ""
                        if "1" in val:
                            lines.append(f"\n# {p_text}\n")
                        elif "2" in val:
                            lines.append(f"\n## {p_text}\n")
                        elif "3" in val:
                            lines.append(f"\n### {p_text}\n")
                        else:
                            lines.append(p_text)

                elif tag == "tbl":
                    # Table
                    rows = child.findall(".//w:tr", ns)
                    if rows:
                        table_data = []
                        for row in rows:
                            cells = row.findall(".//w:tc", ns)
                            row_vals = []
                            for cell in cells:
                                c_texts = [t.text for t in cell.findall(".//w:t", ns) if t.text]
                                row_vals.append(" ".join("".join(c_texts).split()))
                            if row_vals:
                                table_data.append(row_vals)

                        if table_data:
                            lines.append("\n")
                            headers = table_data[0]
                            lines.append("| " + " | ".join(headers) + " |")
                            lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                            for r_vals in table_data[1:]:
                                # Pad or trim if irregular
                                if len(r_vals) < len(headers):
                                    r_vals.extend([""] * (len(headers) - len(r_vals)))
                                lines.append("| " + " | ".join(r_vals[: len(headers)]) + " |")
                            lines.append("\n")

            return "\n".join(lines).strip()
        except Exception as e:
            raise RuntimeError(f"Failed to read .docx file: {str(e)}")

    @staticmethod
    def _read_text(filepath: str) -> str:
        """Read standard plain text or markdown file with multi-encoding fallback."""
        encodings = ["utf-8", "utf-8-sig", "cp1251", "latin-1"]
        for enc in encodings:
            try:
                with open(filepath, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, LookupError):
                continue
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    @staticmethod
    def _read_json(filepath: str) -> str:
        """Read structured JSON file and return pretty-printed string."""
        raw = RequirementsDocReader._read_text(filepath)
        try:
            data = json.loads(raw)
            return json.dumps(data, indent=2, ensure_ascii=False)
        except Exception:
            return raw
